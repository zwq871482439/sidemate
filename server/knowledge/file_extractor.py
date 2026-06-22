# -*- coding: utf-8 -*-
"""
file_extractor.py — 文件内容提取 + 长文件三级策略
支持：txt/md/csv/docx/xlsx/pdf（纯文本提取，图片忽略）
.doc/.xls 旧格式不支持，返回提示让用户转存新格式
"""

import os
import re
import logging

log = logging.getLogger(__name__)

MAX_INLINE_CHARS = 1500      # 直接注入下限（保底值，动态模式下会被覆盖）
MAX_EXTRACT_CHARS = 5000     # 智能截取下限（保底值）

# 动态文件注入预算常量
FILE_BUDGET_TOTAL_TOKENS = 32000     # 总上下文窗口
FILE_BUDGET_SYSTEM_TOKENS = 300      # 系统提示预留
FILE_BUDGET_OUTPUT_TOKENS = 4096     # 输出预留
FILE_BUDGET_HISTORY_RATIO = 0.70     # 文件可用空间占剩余的比例
FILE_BUDGET_MIN_CHARS = 3000         # 保底注入字数
FILE_BUDGET_MAX_CHARS = 25000        # 上限注入字数
CHARS_PER_TOKEN = 1.5                # 粗估字/token 比


def extract_text(file_path: str) -> str:
    """根据扩展名提取文本"""
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext in ('.txt', '.md', '.csv'):
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception as e:
            log.error(f"读取文本文件失败: {e}")
            return ""
    
    elif ext == '.docx':
        try:
            from docx import Document
            doc = Document(file_path)
            return '\n\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            log.warning("python-docx 未安装，无法提取 .docx 文件")
            return ""
        except Exception as e:
            log.error(f"提取 docx 失败: {e}")
            return ""
    
    elif ext == '.doc':
        return "[不支持 .doc 旧格式，请用 Word 另存为 .docx 后重新上传]"
    
    elif ext == '.xlsx':
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True)
            texts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    texts.append(' | '.join(str(c) for c in row if c is not None))
            wb.close()
            return '\n'.join(texts)
        except ImportError:
            log.warning("openpyxl 未安装，无法提取 .xlsx 文件")
            return ""
        except Exception as e:
            log.error(f"提取 xlsx 失败: {e}")
            return ""
    
    elif ext == '.xls':
        return "[不支持 .xls 旧格式，请用 Excel 另存为 .xlsx 后重新上传]"
    
    elif ext == '.pdf':
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(file_path)

            # Patch4 v3.1 BUG#33：加密 PDF 检测
            # fitz.open() 不抛异常，但 page.get_text() 返回空或触发崩溃
            if doc.is_encrypted:
                # 尝试用空密码解锁（很多 PDF 加密但允许空密码阅读）
                if not doc.authenticate(""):
                    doc.close()
                    log.warning("PDF 加密且空密码无法解锁: %s" % file_path)
                    return "[此 PDF 已加密，无法提取文本。请先用 PDF 工具解除密码保护后重新上传]"

            texts = []
            # Patch4 v3.1：100 → 1000，配合 bge-m3 长序列支持大文档
            max_pages = min(len(doc), 1000)
            for page_idx in range(max_pages):
                page = doc[page_idx]
                texts.append(page.get_text())
            doc.close()
            return '\n\n'.join(texts)
        except ImportError:
            # 回退到 pdfplumber
            try:
                import pdfplumber
                with pdfplumber.open(file_path) as pdf:
                    texts = []
                    for page in pdf.pages:
                        t = page.extract_text()
                        if t:
                            texts.append(t)
                    return '\n\n'.join(texts)
            except ImportError:
                log.warning("PyMuPDF 和 pdfplumber 均未安装，无法提取 PDF")
                return ""
        except Exception as e:
            log.error(f"提取 PDF 失败: {e}")
            return ""

    elif ext == '.epub':
        # B2: EPUB 电子书解析（ebooklib + beautifulsoup4）
        try:
            import ebooklib
            from ebooklib import epub
            from bs4 import BeautifulSoup
            book = epub.read_epub(file_path, options={"ignore_ncx": True})
            texts = []
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                html_content = item.get_content().decode("utf-8", errors="replace")
                soup = BeautifulSoup(html_content, "html.parser")
                # 提取章节标题作为 heading
                for tag in soup.find_all(["h1", "h2", "h3"]):
                    title_text = tag.get_text(strip=True)
                    if title_text:
                        texts.append(title_text)
                # 提取段落文本
                for tag in soup.find_all(["p", "div"]):
                    para_text = tag.get_text(strip=True)
                    if para_text:
                        texts.append(para_text)
            result = "\n\n".join(texts)
            if not result.strip():
                return "[此 EPUB 文件无法提取到文本内容，可能为纯图片电子书或 DRM 保护文件]"
            return result
        except ImportError:
            log.warning("ebooklib / beautifulsoup4 未安装，无法提取 .epub 文件")
            return ""
        except Exception as e:
            err_msg = str(e)[:200]
            log.error(f"提取 EPUB 失败: {err_msg}")
            # DRM 保护的 epub 常见错误
            if "drm" in err_msg.lower() or "encrypt" in err_msg.lower():
                return "[此 EPUB 文件受 DRM 保护，无法提取文本。请去除 DRM 后重新上传]"
            return ""

    elif ext in ('.html', '.htm'):
        # B2: HTML 网页文件解析（beautifulsoup4）
        try:
            from bs4 import BeautifulSoup
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                html_content = f.read()
            soup = BeautifulSoup(html_content, "html.parser")
            # 移除 script 和 style 标签
            for tag in soup.find_all(["script", "style", "noscript"]):
                tag.decompose()
            texts = []
            # 提取标题
            title = soup.find("title")
            if title and title.get_text(strip=True):
                texts.append(title.get_text(strip=True))
            # 提取正文内容（p/div/table/li 等）
            for tag in soup.find_all(["p", "div", "li", "td", "th", "h1", "h2", "h3", "h4", "h5", "h6"]):
                tag_text = tag.get_text(strip=True)
                if tag_text and len(tag_text) > 1:
                    texts.append(tag_text)
            return "\n".join(texts)
        except ImportError:
            log.warning("beautifulsoup4 未安装，无法提取 .html 文件")
            return ""
        except Exception as e:
            log.error(f"提取 HTML 失败: {e}")
            return ""

    elif ext == '.srt':
        # B2: SRT 字幕文件解析（纯文本正则，无第三方库）
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            # SRT 格式：序号 + 时间轴 + 字幕文本（空行分隔）
            # 正则去掉序号行和时间轴行，保留字幕文本
            lines = content.splitlines()
            texts = []
            for line in lines:
                line = line.strip()
                # 跳过序号行（纯数字）
                if line.isdigit():
                    continue
                # 跳过时间轴行（包含 --> 格式）
                if "-->" in line:
                    continue
                # 跳过空行
                if not line:
                    continue
                texts.append(line)
            return "\n".join(texts)
        except Exception as e:
            log.error(f"提取 SRT 失败: {e}")
            return ""

    elif ext == '.rtf':
        # B2: RTF 富文本解析（striprtf）
        try:
            from striprtf.striprtf import rtf_to_text
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                rtf_content = f.read()
            text = rtf_to_text(rtf_content)
            return text.strip()
        except ImportError:
            log.warning("striprtf 未安装，无法提取 .rtf 文件")
            return ""
        except Exception as e:
            log.error(f"提取 RTF 失败: {e}")
            return ""
    
    else:
        log.warning(f"不支持的文件类型: {ext}")
        return ""


def _extract_query_keywords(text: str) -> list:
    """
    轻量关键词提取（无需 jieba）。
    中文：2-gram 滑动窗口（"考勤制度" → ["考勤", "勤制", "制度"]）
    英文：按空格分词
    """
    keywords = []
    # 英文词（2+字符）
    en_words = re.findall(r'[a-zA-Z]{2,}', text.lower())
    keywords.extend(en_words)
    # 中文 2-gram
    cn_chars = re.findall(r'[\u4e00-\u9fff]', text)
    for i in range(len(cn_chars) - 1):
        keywords.append(cn_chars[i] + cn_chars[i + 1])
    return keywords


def smart_extract(file_text: str, user_message: str, max_chars: int) -> str:
    """按段落相关性截取，优先保留和用户问题相关的内容"""
    paragraphs = [p.strip() for p in file_text.split('\n\n') if p.strip()]
    
    if not paragraphs:
        return file_text[:max_chars]
    
    # 2-gram 关键词匹配
    query_kws = _extract_query_keywords(user_message)
    if not query_kws:
        # 没有有效关键词，取头尾
        head = file_text[:max_chars // 2]
        tail = file_text[-(max_chars // 2):]
        return f"{head}\n\n...(省略 {len(file_text) - max_chars} 字)...\n\n{tail}"
    
    scored = []
    for para in paragraphs:
        para_lower = para.lower()
        score = sum(1 for kw in query_kws if kw in para_lower)
        scored.append((score, para))
    
    # 按相关性排序
    scored.sort(key=lambda x: -x[0])
    
    # 取前 N 个直到填满（只取有相关性的）
    result = []
    total = 0
    for score, para in scored:
        if score == 0:
            break
        if total + len(para) > max_chars:
            break
        result.append(para)
        total += len(para)
    
    # 如果全都不相关，取头尾
    if not result:
        head = file_text[:max_chars // 2]
        tail = file_text[-(max_chars // 2):]
        return f"{head}\n\n...(省略 {len(file_text) - max_chars} 字)...\n\n{tail}"
    
    return '\n\n'.join(result)


def calc_file_budget(history_chars: int = 0) -> int:
    """
    根据对话历史长度动态计算文件可注入的最大字数。
    
    公式：可用 tokens = (32000 - 300 - 4096 - 历史 tokens) × 0.70
    保底 3000 字，上限 25000 字。
    """
    history_tokens = int(history_chars / CHARS_PER_TOKEN)
    remaining_tokens = FILE_BUDGET_TOTAL_TOKENS - FILE_BUDGET_SYSTEM_TOKENS - FILE_BUDGET_OUTPUT_TOKENS - history_tokens
    file_tokens = int(max(0, remaining_tokens) * FILE_BUDGET_HISTORY_RATIO)
    file_chars = int(file_tokens * CHARS_PER_TOKEN)
    file_chars = max(FILE_BUDGET_MIN_CHARS, min(file_chars, FILE_BUDGET_MAX_CHARS))
    return file_chars


def process_uploaded_file(file_path: str, user_message: str = "", max_chars: int = 0) -> dict:
    """
    处理上传文件，返回结构化结果。
    
    Args:
        max_chars: 文件内容注入上限（动态预算），0 表示使用默认保底值
    
    Returns:
        {
            "status": "ok" | "truncated" | "too_long" | "error",
            "text": str or None,
            "total_chars": int,
            "extracted_chars": int,
            "filename": str,
            "message": str
        }
    """
    filename = os.path.basename(file_path)
    inline_limit = max_chars if max_chars > 0 else MAX_INLINE_CHARS
    extract_limit = inline_limit * 3  # 智能截取上限 = 注入上限的 3 倍

    # 文件大小检查（Patch4 v3.1：50MB → 200MB，允许电子书/长PDF/大型网页存档）
    try:
        file_size = os.path.getsize(file_path)
        if file_size > 200 * 1024 * 1024:
            return {
                "status": "error",
                "text": None,
                "total_chars": 0,
                "extracted_chars": 0,
                "filename": filename,
                "message": " %s 文件过大（%.1fMB），最大支持 200MB" % (filename, file_size / 1024 / 1024),
            }
    except OSError as e:
        log.error(f"文件大小检查失败: {e}")
        return {
            "status": "error",
            "text": None,
            "total_chars": 0,
            "extracted_chars": 0,
            "filename": filename,
            "message": f" 文件访问失败: {e}",
        }

    try:
        text = extract_text(file_path)
    except Exception as e:
        log.error(f"文件提取异常: {e}")
        return {
            "status": "error",
            "text": None,
            "total_chars": 0,
            "extracted_chars": 0,
            "filename": filename,
            "message": f" 文件读取失败: {e}"
        }
    
    total = len(text)
    
    if total == 0:
        return {
            "status": "error",
            "text": "",
            "total_chars": 0,
            "extracted_chars": 0,
            "filename": filename,
            "message": f" {filename} 内容为空或无法提取文本"
        }
    
    if total <= inline_limit:
        # 级别 1：短文件，直接注入
        return {
            "status": "ok",
            "text": text,
            "total_chars": total,
            "extracted_chars": total,
            "filename": filename,
            "message": f" {filename}（{total}字）"
        }
    
    elif total <= extract_limit:
        # 级别 2：中等文件，智能截取
        extracted = smart_extract(text, user_message, inline_limit)
        return {
            "status": "truncated",
            "text": extracted,
            "total_chars": total,
            "extracted_chars": len(extracted),
            "filename": filename,
            "message": f" {filename}（{total}字）→ 已提取与问题最相关的 {len(extracted)} 字\n️ 文件较长，查看完整内容请上传至文库"
        }
    
    else:
        # 级别 3：超长文件，引导文库
        return {
            "status": "too_long",
            "text": None,
            "total_chars": total,
            "extracted_chars": 0,
            "filename": filename,
            "message": f" {filename}（{total}字）\n 文件过长，无法在对话中完整使用\n 建议：上传至文库，可以随时检索提问"
        }
